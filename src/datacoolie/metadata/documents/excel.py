"""Excel workbook parsing for :class:`~datacoolie.metadata.FileProvider`.

This module is intentionally pure with respect to the provider: it receives a
workbook source (or an already opened workbook for sheet helpers) and returns
the same section-wrapper dictionaries used by JSON/YAML metadata.  Physical
Excel row numbers are part of the single row representation so validation
errors always point to the actual source row.
"""

from __future__ import annotations

import io
import json
import math
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

from datacoolie.core.exceptions import MetadataError
from datacoolie.utils.converters import convert_to_bool
from datacoolie.utils.collections import ensure_list


def _convert_excel_integer(value: Any) -> int | None:
    """Parse an integer metadata cell without truncating fractional values."""

    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("Boolean values are not integer metadata")
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        if not math.isfinite(value) or not value.is_integer():
            raise ValueError(f"Expected a finite whole number: {value!r}")
        return int(value)
    if isinstance(value, str):
        stripped = value.strip()
        if not stripped:
            return None
        try:
            number = Decimal(stripped)
        except InvalidOperation as exc:
            raise ValueError(f"Invalid integer metadata: {value!r}") from exc
        if not number.is_finite() or number != number.to_integral_value():
            raise ValueError(f"Expected a finite whole number: {value!r}")
        return int(number)
    raise TypeError(f"Cannot convert to integer metadata: {value!r}")


EXCEL_LIST_COLS: frozenset[str] = frozenset(
    {
        "source_watermark_columns",
        "destination_merge_keys",
        "destination_partition_columns",
        "transform_select_columns",
        "transform_drop_columns",
    }
)
EXCEL_JSON_COLS: frozenset[str] = frozenset(
    {
        "configure",
        "source_configure",
        "destination_configure",
        "transform_deduplicate_columns",
        "transform_latest_data_columns",
        "transform_additional_columns",
        "transform_schema_hints",
        "transform_rename_columns",
        "transform_value_rules",
        "transform_hash_columns",
        "transform_masking_rules",
        "transform_configure",
    }
)


def cast(value: Any) -> Any:
    """Normalize a cell value, representing blank strings as ``None``."""

    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def safe_bool(value: Any) -> bool:
    """Coerce a cell to ``bool``; unrecognized values are inactive."""

    try:
        return convert_to_bool(value)
    except ValueError:
        return False


def json_cell(value: Any) -> Any:
    """Parse an object/array cell; blank and empty JSON values become ``None``."""

    if value is None:
        return None
    if isinstance(value, (dict, list)):
        result = value
    else:
        text = str(value).strip()
        if not text:
            return None
        try:
            result = json.loads(text)
        except (TypeError, json.JSONDecodeError) as exc:
            raise MetadataError(f"Invalid JSON cell value: {value!r}") from exc
    if not isinstance(result, (dict, list)):
        raise MetadataError(f"Invalid JSON cell value: {value!r}")
    return result if result else None


def excel_sheet_rows(wb: Any, sheet_name: str) -> List[Tuple[int, Dict[str, Any]]]:
    """Return nonblank rows as ``(physical_row_number, row_mapping)`` tuples."""

    if sheet_name not in wb.sheetnames:
        return []
    ws = wb[sheet_name]
    header_iter = next(ws.iter_rows(min_row=1, max_row=1), ())
    headers = [
        str(cell.value).strip() if cell.value is not None else ""
        for cell in header_iter
    ]
    rows: List[Tuple[int, Dict[str, Any]]] = []
    for row_number, row in enumerate(
        ws.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if all(value is None for value in row):
            continue
        row_dict = {headers[index]: row[index] for index in range(len(headers))}
        rows.append((row_number, row_dict))
    return rows


def excel_row_context(
    source_path: Optional[str],
    sheet_name: str,
    row_number: int,
) -> str:
    """Build a stable file/sheet/row prefix for Excel parse errors."""

    return f"{source_path or '<workbook>'} [{sheet_name} row {row_number}]"


def parse_excel(
    source: str | bytes,
    *,
    source_path: Optional[str] = None,
) -> Dict[str, Any]:
    """Parse an Excel workbook into standard metadata section wrappers.

    Individual sheets are optional, but a workbook must contain at least one
    of the supported metadata sheets; arbitrary workbooks are rejected before
    they can be interpreted as an empty metadata shard.
    """

    try:
        import openpyxl  # noqa: WPS433 — optional dependency
    except ImportError as exc:
        raise MetadataError(
            "openpyxl is required for Excel metadata files.  "
            "Install it with:  pip install openpyxl"
        ) from exc
    display_path = source_path or (source if isinstance(source, str) else "<bytes>")
    try:
        workbook_source: Any = io.BytesIO(source) if isinstance(source, bytes) else source
        workbook = openpyxl.load_workbook(
            workbook_source,
            read_only=True,
            data_only=True,
        )
    except Exception as exc:
        raise MetadataError(f"Cannot read Excel metadata file: {display_path}") from exc
    try:
        expected_sheets = {"connections", "dataflows", "schema_hints"}
        if not expected_sheets.intersection(set(workbook.sheetnames)):
            raise MetadataError(
                f"Excel metadata file contains none of the supported sheets: {display_path}"
            )
        sheets = set(workbook.sheetnames)
        sections: Dict[str, Any] = {}
        if "connections" in sheets:
            sections["connections"] = parse_excel_connections(
                workbook,
                source_path=source_path,
            )
        if "dataflows" in sheets:
            sections["dataflows"] = parse_excel_dataflows(
                workbook,
                source_path=source_path,
            )
        if "schema_hints" in sheets:
            sections["schema_hints"] = parse_excel_schema_hints(
                workbook,
                source_path=source_path,
            )
        return sections
    finally:
        workbook.close()


def parse_excel_connections(
    wb: Any,
    *,
    source_path: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Parse the ``connections`` worksheet."""

    connections: List[Dict[str, Any]] = []
    for row_number, row in excel_sheet_rows(wb, "connections"):
        if not any(cast(value) is not None for value in row.values()):
            continue
        name = cast(row.get("name"))
        connection_type = cast(row.get("connection_type"))
        context = excel_row_context(source_path, "connections", row_number)
        if not name or not connection_type:
            raise MetadataError(
                f"Invalid Excel metadata row {context}: "
                "name and connection_type are required"
            )
        connection: Dict[str, Any] = {}
        configure: Dict[str, Any] = {}
        try:
            for key, raw_value in row.items():
                value = cast(raw_value)
                if key == "configure":
                    parsed = json_cell(raw_value)
                    if isinstance(parsed, dict):
                        configure.update(parsed)
                elif key.startswith("configure_"):
                    sub_key = key[len("configure_"):]
                    if value is not None:
                        configure[sub_key] = value
                elif key == "secrets_ref":
                    parsed = json_cell(raw_value)
                    if parsed is not None:
                        connection[key] = parsed
                elif key == "is_active":
                    if value is not None:
                        connection[key] = safe_bool(value)
                elif value is not None:
                    connection[key] = value
        except MetadataError as exc:
            raise MetadataError(f"Invalid Excel metadata row {context}: {exc}") from exc
        if configure:
            connection["configure"] = configure
        if connection:
            connections.append(connection)
    return connections


def parse_excel_dataflows(
    wb: Any,
    *,
    source_path: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Parse the ``dataflows`` worksheet."""

    dataflows: List[Dict[str, Any]] = []
    for row_number, row in excel_sheet_rows(wb, "dataflows"):
        if not any(cast(value) is not None for value in row.values()):
            continue
        source_connection = cast(
            row.get("source_connection_name") or row.get("source_connection")
        )
        source_table = cast(row.get("source_table"))
        source_query = cast(row.get("source_query"))
        source_function = cast(row.get("source_python_function"))
        destination_connection = cast(
            row.get("destination_connection_name")
            or row.get("destination_connection")
        )
        destination_table = cast(row.get("destination_table"))
        context = excel_row_context(source_path, "dataflows", row_number)
        if not source_connection or not (source_table or source_query or source_function):
            raise MetadataError(
                f"Invalid Excel metadata row {context}: source connection and "
                "one of source_table, source_query, or source_python_function are required"
            )
        if not destination_connection or not destination_table:
            raise MetadataError(
                f"Invalid Excel metadata row {context}: "
                "destination connection and destination_table are required"
            )
        dataflow: Dict[str, Any] = {}
        source: Dict[str, Any] = {}
        destination: Dict[str, Any] = {}
        transform: Dict[str, Any] = {}
        transform_base: Dict[str, Any] = {}
        try:
            for key, raw_value in row.items():
                if key == "transform":
                    parsed = json_cell(raw_value)
                    if isinstance(parsed, dict):
                        transform_base = parsed
                    continue
                if key == "is_active":
                    parsed_is_active = cast(raw_value)
                    if parsed_is_active is not None:
                        dataflow["is_active"] = safe_bool(parsed_is_active)
                    continue
                if key.startswith("source_"):
                    target, sub_key = source, key[len("source_"):]
                elif key.startswith("destination_"):
                    target, sub_key = destination, key[len("destination_"):]
                elif key.startswith("transform_"):
                    target, sub_key = transform, key[len("transform_"):]
                else:
                    value = json_cell(raw_value) if key in EXCEL_JSON_COLS else cast(raw_value)
                    if value is not None:
                        dataflow[key] = value
                    continue
                if key in EXCEL_LIST_COLS:
                    parsed_list = ensure_list(raw_value)
                    if parsed_list:
                        target[sub_key] = parsed_list
                elif key in EXCEL_JSON_COLS:
                    parsed_json = json_cell(raw_value)
                    if parsed_json is not None:
                        target[sub_key] = parsed_json
                else:
                    value = cast(raw_value)
                    if value is not None:
                        target[sub_key] = value
        except MetadataError as exc:
            raise MetadataError(f"Invalid Excel metadata row {context}: {exc}") from exc
        if source:
            dataflow["source"] = source
        if destination:
            dataflow["destination"] = destination
        merged_transform = {**transform_base, **transform}
        if merged_transform:
            dataflow["transform"] = merged_transform
        if dataflow:
            dataflows.append(dataflow)
    return dataflows


def parse_excel_schema_hints(
    wb: Any,
    *,
    source_path: Optional[str] = None,
) -> List[Dict[str, Any]]:
    """Parse and group rows from the ``schema_hints`` worksheet."""

    hints_map: Dict[tuple, Dict[str, Any]] = {}
    for row_number, row in excel_sheet_rows(wb, "schema_hints"):
        connection_name = cast(row.get("connection_name"))
        connection_id = cast(row.get("connection_id"))
        table_name = cast(row.get("table_name"))
        schema_name = cast(row.get("schema_name"))
        connection_ref = connection_name or connection_id
        if not connection_ref or not table_name:
            if any(cast(value) is not None for value in row.values()):
                raise MetadataError(
                    "Invalid schema_hints row "
                    f"{excel_row_context(source_path, 'schema_hints', row_number)}: "
                    "connection_name and table_name (or connection_id and table_name) are required"
                )
            continue
        column_name = cast(row.get("column_name"))
        data_type = cast(row.get("data_type"))
        if not column_name or not data_type:
            raise MetadataError(
                "Invalid schema_hints row "
                f"{excel_row_context(source_path, 'schema_hints', row_number)}: "
                "column_name and data_type are required"
            )
        # Group by the stable identity when it is authored, while retaining
        # the display name if the workbook contains both reference forms.
        group_key = (
            f"id:{connection_id}" if connection_id else f"name:{connection_name}",
            table_name,
            schema_name,
        )
        if group_key not in hints_map:
            hints_map[group_key] = {
                "table_name": table_name,
                "hints": [],
            }
            if connection_name is not None:
                hints_map[group_key]["connection_name"] = connection_name
            if connection_id is not None:
                hints_map[group_key]["connection_id"] = connection_id
            if schema_name is not None:
                hints_map[group_key]["schema_name"] = schema_name
        else:
            if connection_name is not None:
                hints_map[group_key].setdefault("connection_name", connection_name)
            if connection_id is not None:
                hints_map[group_key].setdefault("connection_id", connection_id)
        hint: Dict[str, Any] = {}
        for field in ("column_name", "data_type", "format"):
            value = cast(row.get(field))
            if value is not None:
                hint[field] = value
        for field in ("precision", "scale"):
            try:
                converted = _convert_excel_integer(row.get(field))
            except (TypeError, ValueError) as exc:
                raise MetadataError(
                    "Invalid schema_hints cell "
                    f"{excel_row_context(source_path, 'schema_hints', row_number)} "
                    f"field {field!r}: {row.get(field)!r}"
                ) from exc
            if converted is not None:
                hint[field] = converted
        default_value = cast(row.get("default_value"))
        if default_value is not None:
            hint["default_value"] = default_value
        try:
            ordinal_position = _convert_excel_integer(row.get("ordinal_position"))
        except (TypeError, ValueError) as exc:
            raise MetadataError(
                "Invalid schema_hints cell "
                f"{excel_row_context(source_path, 'schema_hints', row_number)} "
                "field 'ordinal_position': "
                f"{row.get('ordinal_position')!r}"
            ) from exc
        if ordinal_position is not None:
            hint["ordinal_position"] = ordinal_position
        parsed_is_active = cast(row.get("is_active"))
        if parsed_is_active is not None:
            hint["is_active"] = safe_bool(parsed_is_active)
        if hint:
            hints_map[group_key]["hints"].append(hint)
    return list(hints_map.values())


__all__ = [
    "EXCEL_JSON_COLS",
    "EXCEL_LIST_COLS",
    "cast",
    "excel_row_context",
    "excel_sheet_rows",
    "json_cell",
    "parse_excel",
    "parse_excel_connections",
    "parse_excel_dataflows",
    "parse_excel_schema_hints",
    "safe_bool",
]
