"""Shared project metadata document loading, inventory, and serialization."""

from __future__ import annotations

from dataclasses import dataclass, field
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
from typing import Any, Iterable

from datacoolie.metadata.documents.parsers import (
    SUPPORTED_METADATA_SUFFIXES,
    parse_json_document,
    parse_yaml_document,
    validate_document,
)
from datacoolie.metadata.documents.excel import parse_excel

from .errors import ProjectDependencyError, ProjectValidationError


SECTION_KEYS = ("connections", "dataflows", "schema_hints")
_IGNORED_NAMES = {".gitkeep", ".DS_Store", "Thumbs.db"}
_FORMAT_TO_SUFFIX = {"json": ".json", "yaml": ".yml", "excel": ".xlsx"}
_SUFFIX_TO_FORMAT = {
    ".json": "json",
    ".yaml": "yaml",
    ".yml": "yaml",
    ".xlsx": "excel",
}


def format_for_path(path: Path | str) -> str:
    suffix = Path(path).suffix.lower()
    try:
        return _SUFFIX_TO_FORMAT[suffix]
    except KeyError as exc:
        raise ProjectValidationError(
            f"Unsupported metadata format {suffix or '<none>'!r}; "
            "expected .json, .yaml, .yml, or .xlsx"
        ) from exc


def _copy_document(value: Any) -> Any:
    # Avoid exposing mutable parser internals to overlay/build code without
    # requiring pydantic; codecs can still report unsupported non-JSON values
    # at serialization time.
    return deepcopy(value)


def decode_bytes(raw: bytes, *, path: Path | str) -> dict[str, Any]:
    """Decode one supported metadata file into section-wrapper mappings."""

    display = str(path)
    fmt = format_for_path(path)
    try:
        if fmt == "yaml":
            __import__("yaml")
        elif fmt == "excel":
            __import__("openpyxl")
    except ImportError as exc:
        if fmt == "yaml":
            raise ProjectDependencyError(
                "PyYAML is required for YAML metadata files; install datacoolie[cli]",
                dependency="PyYAML",
            ) from exc
        raise ProjectDependencyError(
            "openpyxl is required for Excel metadata files; install datacoolie[metadata-excel]",
            dependency="openpyxl",
        ) from exc
    try:
        if fmt == "excel":
            document = parse_excel(raw, source_path=display)
        elif fmt == "yaml":
            document = parse_yaml_document(raw.decode("utf-8"), display)
        else:
            document = parse_json_document(raw.decode("utf-8"), display)
        return validate_document(document, display)
    except Exception as exc:
        if isinstance(exc, (ProjectDependencyError, ProjectValidationError)):
            raise
        raise ProjectValidationError(f"Cannot decode metadata document: {display}") from exc


def decode_file(path: Path) -> dict[str, Any]:
    try:
        raw = path.read_bytes()
    except OSError as exc:
        raise ProjectValidationError(f"Cannot read metadata document: {path}") from exc
    return decode_bytes(raw, path=path)


@dataclass
class MetadataDocument:
    relative_path: str
    format: str
    payload: dict[str, Any]
    sections: tuple[str, ...]


@dataclass
class MetadataSnapshot:
    """Raw metadata plus document/entity provenance for project operations."""

    root: Path
    documents: list[MetadataDocument] = field(default_factory=list)
    # Top-level values which are not one of the framework sections (for
    # example ``$schema`` or a project-owned extension).  Keeping these values
    # here prevents aggregate build layouts from silently dropping authored
    # content.  Conflicting values across shards are rejected during loading;
    # there is no principled merge rule for arbitrary extension fields.
    extras: dict[str, Any] = field(default_factory=dict)
    sections: dict[str, list[dict[str, Any]]] = field(
        default_factory=lambda: {key: [] for key in SECTION_KEYS}
    )
    origins: dict[tuple[str, str], str] = field(default_factory=dict)

    @property
    def is_empty(self) -> bool:
        return not self.documents or not any(self.sections.values())

    def merged(self) -> dict[str, Any]:
        result = _copy_document(self.extras)
        for section in SECTION_KEYS:
            if self.sections[section]:
                result[section] = _copy_document(self.sections[section])
            else:
                result[section] = []
        return result


def _source_paths(root: Path) -> Iterable[Path]:
    if root.is_file():
        yield root
        return
    if not root.exists():
        raise ProjectValidationError(f"Metadata path does not exist: {root}")
    for path in sorted(root.rglob("*"), key=lambda item: item.as_posix().casefold()):
        if path.is_symlink():
            raise ProjectValidationError(f"Metadata path must not contain symlinks: {path}")
        if not path.is_file() or path.name in _IGNORED_NAMES:
            continue
        relative = path.relative_to(root)
        if any(part.casefold() == "environments" for part in relative.parts):
            continue
        if path.suffix.lower() in SUPPORTED_METADATA_SUFFIXES:
            yield path


def _identity(section: str, item: dict[str, Any], index: int) -> str:
    if section in {"connections", "dataflows"}:
        return str(item.get("name") or item.get("dataflow_id") or item.get("connection_id") or f"#{index}")
    connection = item.get("connection_name") or item.get("connection_id") or "?"
    schema = item.get("schema_name") or ""
    table = item.get("table_name") or "?"
    return f"{connection}|{schema}|{table}"


def load_snapshot(root: Path | str) -> MetadataSnapshot:
    """Load all supported files below *root* in deterministic order.

    Section names come from document content, never from a filename.  The
    environment subtree is deliberately excluded because it is an overlay,
    not ordinary common metadata.
    """

    candidate = Path(root).expanduser()
    if candidate.is_symlink():
        raise ProjectValidationError(f"Metadata path must not be a symlink: {candidate}")
    selected = candidate.resolve()
    snapshot = MetadataSnapshot(root=selected)
    if not selected.exists():
        raise ProjectValidationError(f"Metadata path does not exist: {selected}")
    for path in _source_paths(selected):
        payload = decode_file(path)
        relative = path.name if selected.is_file() else path.relative_to(selected).as_posix()
        sections = tuple(section for section in SECTION_KEYS if section in payload)
        document = MetadataDocument(relative, format_for_path(path), _copy_document(payload), sections)
        snapshot.documents.append(document)
        for key, value in payload.items():
            if key in SECTION_KEYS:
                continue
            if key in snapshot.extras and canonical_json(snapshot.extras[key]) != canonical_json(value):
                raise ProjectValidationError(
                    f"Conflicting top-level metadata field {key!r}: "
                    f"{relative} differs from an earlier document"
                )
            snapshot.extras.setdefault(key, _copy_document(value))
        for section in sections:
            values = payload[section]
            if not isinstance(values, list):
                raise ProjectValidationError(
                    f"Metadata section {section!r} must be a list: {path}"
                )
            for index, item in enumerate(values):
                if not isinstance(item, dict):
                    raise ProjectValidationError(
                        f"Metadata section {section!r} item {index} must be an object: {path}"
                    )
                copied = _copy_document(item)
                snapshot.sections[section].append(copied)
                snapshot.origins[(section, _identity(section, copied, index))] = relative
    return snapshot


def _json_value(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return value


def _excel_rows(section: str, values: list[dict[str, Any]]) -> tuple[list[str], list[list[Any]]]:
    if section == "connections":
        rows: list[dict[str, Any]] = []
        for item in values:
            row = {key: value for key, value in item.items() if key != "configure"}
            if isinstance(item.get("configure"), dict):
                # The FileProvider Excel parser treats ``configure`` as one
                # JSON cell.  Keeping nested values in that cell avoids
                # turning objects such as read_options into strings.
                row["configure"] = _json_value(item["configure"])
            rows.append(row)
    elif section == "dataflows":
        rows = []
        for item in values:
            row: dict[str, Any] = {}
            for key, value in item.items():
                if key in {"source", "destination", "transform"}:
                    continue
                row[key] = _json_value(value)
            for prefix in ("source", "destination", "transform"):
                nested = item.get(prefix)
                if not isinstance(nested, dict):
                    continue
                for key, value in nested.items():
                    row[f"{prefix}_{key}"] = _json_value(value)
            rows.append(row)
    else:
        rows = []
        for group in values:
            group_fields = {key: value for key, value in group.items() if key != "hints"}
            for hint in group.get("hints", []) if isinstance(group.get("hints"), list) else []:
                rows.append({**group_fields, **hint})
    headers: list[str] = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    if section == "connections":
        headers = [key for key in ("name", "connection_id", "workspace_id", "connection_type", "format", "catalog", "database", "is_active", "secrets_ref") if key in headers] + [key for key in headers if key not in {"name", "connection_id", "workspace_id", "connection_type", "format", "catalog", "database", "is_active", "secrets_ref"}]
    elif section == "dataflows":
        preferred = ["name", "dataflow_id", "stage", "description", "is_active", "source_connection_name", "source_table", "source_query", "source_python_function", "destination_connection_name", "destination_table"]
        headers = [key for key in preferred if key in headers] + [key for key in headers if key not in preferred]
    elif section == "schema_hints":
        preferred = ["connection_name", "connection_id", "table_name", "schema_name", "column_name", "data_type", "precision", "scale", "format", "default_value", "ordinal_position", "is_active"]
        headers = [key for key in preferred if key in headers] + [key for key in headers if key not in preferred]
    if not headers:
        headers = {
            "connections": ["name", "connection_type"],
            "dataflows": ["name", "source_connection_name", "source_table", "destination_connection_name", "destination_table"],
            "schema_hints": ["connection_name", "table_name", "column_name", "data_type"],
        }[section]
    return headers, [[row.get(header) for header in headers] for row in rows]


def encode_document(document: dict[str, Any], fmt: str) -> bytes:
    """Encode a section-wrapper document as JSON, YAML, or XLSX bytes."""

    if fmt == "json":
        return (json.dumps(document, indent=2, ensure_ascii=False) + "\n").encode("utf-8")
    if fmt == "yaml":
        try:
            import yaml  # noqa: WPS433 - optional CLI dependency
        except ImportError as exc:
            raise ProjectDependencyError(
                "PyYAML is required for YAML output; install datacoolie[cli]",
                dependency="PyYAML",
            ) from exc
        return yaml.safe_dump(document, sort_keys=False, allow_unicode=True).encode("utf-8")
    if fmt != "excel":
        raise ProjectValidationError(f"Unsupported metadata output format: {fmt}")
    try:
        import openpyxl  # noqa: WPS433 - optional metadata dependency
    except ImportError as exc:
        raise ProjectDependencyError(
            "openpyxl is required for Excel output; install datacoolie[metadata-excel]",
            dependency="openpyxl",
        ) from exc
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)
    for section in SECTION_KEYS:
        values = document.get(section)
        if not isinstance(values, list):
            continue
        worksheet = workbook.create_sheet(section)
        headers, rows = _excel_rows(section, values)
        worksheet.append(headers)
        for row in rows:
            worksheet.append(row)
    if not workbook.sheetnames:
        workbook.create_sheet("connections").append(["name", "connection_type"])
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def check_output_format(fmt: str) -> None:
    """Check optional dependencies needed to encode ``fmt`` without writing.

    Build planning uses this small preflight so a dry-run reports an absent
    codec before normal execution reaches the staging directory.  The actual
    encoder still performs its own import and remains the final authority.
    """

    if fmt in {"json", "preserve"}:
        return
    if fmt == "yaml":
        module_name = "yaml"
        dependency = "PyYAML"
        install_hint = "datacoolie[cli]"
    elif fmt == "excel":
        module_name = "openpyxl"
        dependency = "openpyxl"
        install_hint = "datacoolie[metadata-excel]"
    else:
        raise ProjectValidationError(f"Unsupported metadata output format: {fmt}")
    try:
        __import__(module_name)
    except ImportError as exc:
        raise ProjectDependencyError(
            f"{dependency} is required for {fmt.upper()} output; install {install_hint}",
            dependency=dependency,
        ) from exc


def write_document(path: Path, document: dict[str, Any], fmt: str) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    try:
        path.write_bytes(encode_document(document, fmt))
    except OSError as exc:
        raise ProjectValidationError(f"Cannot write metadata document: {path}") from exc


def canonical_json(value: Any) -> str:
    return json.dumps(_semantic_value(value), sort_keys=True, ensure_ascii=False, separators=(",", ":"))


def _semantic_value(value: Any) -> Any:
    """Normalize representation-only differences for conversion checks.

    FileProvider models treat omitted nullable fields and explicit ``null`` as
    the same value (notably ``schema_name``).  Excel has no faithful null-cell
    distinction, so only ``None`` is removed here; false/zero/empty strings
    remain significant.
    """

    if isinstance(value, dict):
        return {
            str(key): _semantic_value(item)
            for key, item in value.items()
            if item is not None
        }
    if isinstance(value, list):
        return [_semantic_value(item) for item in value]
    return value


def output_suffix(fmt: str) -> str:
    try:
        return _FORMAT_TO_SUFFIX[fmt]
    except KeyError as exc:
        raise ProjectValidationError(f"Unsupported metadata output format: {fmt}") from exc


__all__ = [
    "MetadataDocument",
    "MetadataSnapshot",
    "SECTION_KEYS",
    "canonical_json",
    "check_output_format",
    "decode_bytes",
    "decode_file",
    "encode_document",
    "format_for_path",
    "load_snapshot",
    "output_suffix",
    "write_document",
]
